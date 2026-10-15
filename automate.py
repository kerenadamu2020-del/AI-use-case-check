"""
automate.py - ask an AI to draft a first-pass risk assessment for each use case in
batch_use_cases.xlsx, then write the draft back into the sheet.

The AI only DRAFTS. A human reviews every row (columns H-K) and decides.
Use FICTIONAL or organisation-approved data only.

Modes (set the PROVIDER environment variable):
  ollama     (default) free, runs on your own computer, no key, no data leaves your machine
             Setup: install Ollama, then run:  ollama pull llama3.2
  anthropic  uses the Anthropic API; needs ANTHROPIC_API_KEY (paid account)
             Never put your key in the code or on GitHub.

Run:
  pip install openpyxl requests
  python automate.py
"""
import json, os, re, sys
import requests
from openpyxl import load_workbook

FILE = "batch_use_cases.xlsx"
PROVIDER = os.getenv("PROVIDER", "ollama").lower()
MODEL = os.getenv("MODEL", "llama3.2" if PROVIDER == "ollama" else "claude-sonnet-5-5")

PROMPT = """You are a cautious GRC analyst reviewing an AI use case before approval.
Return ONLY a JSON object, with no other text, using exactly these keys:
  "risk_level": "Low" or "Medium" or "High",
  "main_risks": a short string naming the 2-4 main risks (e.g. hallucination, data leakage, over-reliance, excessive autonomy, legal),
  "human_oversight": a short string saying what human review is needed,
  "recommendation": "Approve" or "Approve with conditions" or "Do not approve yet".
If you are unsure, say so inside main_risks. Be concise.

Use case: {use_case}
Data involved: {data}
"""

def ask(prompt):
    if PROVIDER == "ollama":
        r = requests.post("http://localhost:11434/api/chat", timeout=300, json={
            "model": MODEL, "stream": False,
            "messages": [{"role": "user", "content": prompt}]})
        r.raise_for_status()
        return r.json()["message"]["content"]
    if PROVIDER == "anthropic":
        key = os.getenv("ANTHROPIC_API_KEY")
        if not key:
            sys.exit("Set the ANTHROPIC_API_KEY environment variable first.")
        r = requests.post("https://api.anthropic.com/v1/messages", timeout=120,
            headers={"x-api-key": key, "anthropic-version": "2023-06-01", "content-type": "application/json"},
            json={"model": MODEL, "max_tokens": 500, "messages": [{"role": "user", "content": prompt}]})
        r.raise_for_status()
        return r.json()["content"][0]["text"]
    sys.exit("PROVIDER must be 'ollama' or 'anthropic'.")

def parse(text):
    text = re.sub(r"```(?:json)?", "", text).strip()
    m = re.search(r"\{.*\}", text, re.S)
    return json.loads(m.group(0)) if m else None

wb = load_workbook(FILE)
ws = wb["Batch"]
for row in range(2, ws.max_row + 1):
    uc, data = ws.cell(row, 2).value, ws.cell(row, 3).value
    if not uc or ws.cell(row, 4).value:        # skip empty rows and rows already filled
        continue
    print(f"Row {row}: {uc[:60]}...")
    try:
        out = parse(ask(PROMPT.format(use_case=uc, data=data)))
    except Exception as e:
        print("  failed:", e); continue
    if not out:
        print("  could not read the answer; fill this row by hand"); continue
    ws.cell(row, 4, out.get("risk_level"))
    ws.cell(row, 5, out.get("main_risks"))
    ws.cell(row, 6, out.get("human_oversight"))
    ws.cell(row, 7, out.get("recommendation"))
wb.save(FILE)
print("Done. Now review EVERY row: fill columns H-K and correct what the AI got wrong.")
